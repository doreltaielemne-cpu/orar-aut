"""Build the timetable website from the official UTCN Automatica workbook (.xlsx).

usage: python scripts/build_site.py ORAR.xlsx dist/ [--previous dist/data.json]
Stops with an error (and builds nothing) if the workbook looks broken, so a bad
download never replaces a good site.
"""
import sys, os, re, json, shutil, hashlib, argparse, collections
import openpyxl
sys.path.insert(0, os.path.dirname(__file__))
import xparse, sitedata

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHEETS = [  # sheet name, internal page id, labels
    ('AN I', 3, 'Anul I', 'română', 'Anul I · Automatică și Informatică Aplicată'),
    ('AN I En', 4, 'Anul I', 'engleză', 'Year I · AIA English'),
    ('AN II', 5, 'Anul II', 'română', 'Anul II · Automatică și Informatică Aplicată'),
    ('AN II En', 6, 'Anul II', 'engleză', 'Year II · AIA English'),
    ('AN III', 7, 'Anul III', 'română', 'Anul III · Automatică și Informatică Aplicată'),
    ('AN III En', 8, 'Anul III', 'engleză', 'Year III · AIA English'),
    ('AN IV ', 9, 'Anul IV', 'română', 'Anul IV · Automatică și Informatică Aplicată'),
    ('AN IV En', 10, 'Anul IV', 'engleză', 'Year IV · AIA English'),
    ('MASTER AN I', 11, 'Master anul I', '', 'Master anul I'),
    ('MASTER AN II', 12, 'Master anul II', '', 'Master anul II'),
]
SUB_NORM = [
    (r'SERIA\s+I\s*-\s*A', 'seria I-A'), (r'SERIA\s+II\s*-\s*B', 'seria II-B'),
    (r'^Automatica\s*-\s*A', 'Automatică'), (r'^Informatica\s+aplicat[aă]\s*-\s*IA', 'Informatică aplicată'),
    (r'Cont[r]?olul Avansat al Proceselor', 'Controlul avansat al proceselor (CAP)'),
    (r'Inf\. Apl\. în Ing\. Sist\. Complexe', 'Inf. aplicată în ing. sist. complexe (IAISC)'),
    (r'^\(?IAISC\)?$', 'IAISC'), (r'^\(?CAP\)?$', 'Controlul avansat al proceselor (CAP)'), (r'^\(?ICAF\)?$', 'ICAF'),
    (r'Informatica Aplicata', 'Informatică aplicată (IA)'), (r'^CPS$', 'CPS'),
]


def norm_sub(t):
    t = ' '.join(str(t).split())
    for rx, out in SUB_NORM:
        if re.search(rx, t, re.I):
            return out
    return t


def read_updated(wb):
    ws = wb['ANUNT']
    for row in ws.iter_rows(values_only=True):
        for v in row:
            if isinstance(v, str):
                m = re.search(r'(\d{2}\.\d{2}\.\d{4}),?\s*ora\s*(\d{1,2}:\d{2})', v)
                if m:
                    return f'{m.group(1)}, ora {m.group(2)}'
    return None


def read_footnotes(wb):
    """'Laborator **CP saptamanile 3(4h,18-22), ...' -> {'**CP': '3(4h,18-22), ...'}"""
    fn = {}
    for name in ('MASTER AN I', 'MASTER AN II'):
        ws = wb[name]
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if isinstance(v, str):
                    m = re.search(r'Laborator\s+(\*+)\s*([A-Za-z]+)\s+s[aă]pt[a-z]*\s+(.+)', v)
                    if m:
                        fn[m.group(1) + m.group(2)] = ' '.join(m.group(3).split())
    return fn


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('xlsx'); ap.add_argument('out'); ap.add_argument('--previous')
    a = ap.parse_args()
    wb = openpyxl.load_workbook(a.xlsx)
    cfg = json.load(open(os.path.join(ROOT, 'config.json'), encoding='utf-8'))
    updated = read_updated(wb)
    footnotes = read_footnotes(wb)
    E = []
    group_meta = collections.OrderedDict()
    problems = []
    for sheet, pn, label, line, title in SHEETS:
        if sheet not in wb.sheetnames:
            problems.append(f'lipsește foaia {sheet!r}'); continue
        ws = wb[sheet]
        res = xparse.parse_sheet(ws, sheet.strip())
        S = xparse.Sheet(ws)
        if len(res['groups']) < 2 or len(res['entries']) < 15:
            problems.append(f'{sheet}: {len(res["groups"])} grupe, {len(res["entries"])} activități')
        for g in res['groups']:
            code = g['code']
            sub = ''
            # heading above the group (seria / specialisation), else the text in brackets
            above = S.val(g['row'] - 1, g['cols'][0]) if g['row'] > 1 else ''
            if above and not re.search(r'^(MASTER|ZI|ORA|DAY|TIME)|YEAR|ORAR|UNIVERSITAR|ACADEMIC', above, re.I):
                sub = norm_sub(above)
            m = re.search(r'\(([A-Z]{1,4})\)', g['label'])
            if not sub and m:
                sub = {'A': 'Automatică', 'I': 'Informatică aplicată', 'ISA': 'ISA'}.get(m.group(1), m.group(1))
            if pn == 10 and m:
                sub = {'A': 'Automation', 'I': 'Applied Informatics'}.get(m.group(1), sub)
            sgs = sorted({n for c_, n, _ in res['sgcols'] if c_ == code})
            gid = f"{pn}-{code.replace('#', '_')}"
            group_meta[gid] = dict(id=gid, page=pn, code=code.split('#')[0],
                                   label=code if '#' not in code else code.split('#')[0] + ' (a doua coloană)',
                                   sub=sub, sgs=sgs, entries=[])
        for e in res['entries']:
            t = e['text']
            # expand lab footnote stars (*TAS, **CP ...) with their weeks from the legend
            m = re.match(r'^\s*(\*+)\s*([A-Za-z]+)', t)
            if m and (m.group(1) + m.group(2)) in footnotes:
                t = t + ' [săpt. ' + footnotes[m.group(1) + m.group(2)] + ']'
            E.append(dict(page=pn, day=e['day'], s=e['s'], e=e['e'], parity=e['parity'], cols=e['cols'], text=t))
    # overrides for known quirks of the official file (applied only while the cell text still matches)
    for ov in cfg.get('overrides', []):
        for e in E:
            if e['page'] == ov['page'] and re.search(ov['text'], e['text']) and e['day'] == ov['day']:
                e.update(ov['set'])
    if problems:
        sys.exit('Orarul pare stricat, nu public nimic:\n  ' + '\n  '.join(problems))
    total = len(E)
    if a.previous and os.path.exists(a.previous):
        prev = json.load(open(a.previous, encoding='utf-8'))
        ptot = prev.get('meta', {}).get('rawEntries', 0)
        if ptot and total < 0.6 * ptot:
            sys.exit(f'Prea puține activități față de versiunea precedentă ({total} față de {ptot}); nu public nimic.')
    sitedata.learn_names(E)
    for e in E:
        d = sitedata.describe(e)
        wk = sitedata.weeks_of(e['text'], e['parity'])
        s_, e_ = e['s'], e['e']
        item = dict(d=e['day'], s=int(s_) if s_ == int(s_) else s_, e=int(e_) if e_ == int(e_) else e_, p=e['parity'], w=wk,
                    raw=sitedata.clean(e['text']), **{k: v for k, v in d.items() if v})
        by = collections.defaultdict(list)
        for g, sg in e['cols']:
            by[g].append(int(sg))
        for g, sgs in by.items():
            gid = f"{e['page']}-{g.replace('#', '_')}"
            if gid in group_meta:
                gm = group_meta[gid]
                gm['entries'].append(dict(item, sg=sorted(sgs), all=sorted(sgs) == gm['sgs'], n=len(e['cols'])))
    groups = [g for g in group_meta.values() if g['entries']]
    for g in groups:
        g['entries'].sort(key=lambda x: (x['d'], x['s'], x['e'], x['sg']))
    pages = {str(pn): dict(label=label, line=line, title=title, notes=cfg.get('notes', {}).get(sheet.strip(), []))
             for sheet, pn, label, line, title in SHEETS}
    data = dict(meta=dict(updated=updated, rawEntries=total, source=cfg.get('source_url'), skin=cfg.get('skin', 'grafit')), weeks=cfg['weeks'], holidays=cfg.get('holidays', []),
                pages=pages, groups=groups)
    # ---- write site
    os.makedirs(a.out, exist_ok=True)
    tpl = open(os.path.join(ROOT, 'site', 'template.html'), encoding='utf-8').read()
    blob = json.dumps(data, ensure_ascii=False, separators=(',', ':')).replace('</', '<\\/')
    html = tpl.replace('__DATA__', blob)
    open(os.path.join(a.out, 'index.html'), 'w', encoding='utf-8').write(html)
    json.dump(data, open(os.path.join(a.out, 'data.json'), 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))
    ver = hashlib.sha1(html.encode()).hexdigest()[:10]
    sw = open(os.path.join(ROOT, 'site', 'sw.template.js'), encoding='utf-8').read().replace('__VERSION__', ver)
    open(os.path.join(a.out, 'sw.js'), 'w', encoding='utf-8').write(sw)
    for f in ('manifest.webmanifest', '_headers'):
        shutil.copy(os.path.join(ROOT, 'site', f), os.path.join(a.out, f))
    shutil.copytree(os.path.join(ROOT, 'site', 'icons'), os.path.join(a.out, 'icons'), dirs_exist_ok=True)
    print(f'OK: {len(groups)} grupe, {total} activități, actualizat {updated}, versiune {ver}')


if __name__ == '__main__':
    main()
